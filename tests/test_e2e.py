"""End to end: register, save token, pick a commit, run the analysis, open a fix MR."""
import re
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

FIXTURE = Path(__file__).parent / "fixture_repo"


@pytest.fixture(scope="module")
def client(fake_services, monkeypatch_module):
    from app import gitlab_client
    monkeypatch_module.setattr(gitlab_client.GitLabClient, "checkout",
                               lambda self, url, sha, branch, dest: shutil.copytree(FIXTURE, dest))
    from app.main import app
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def monkeypatch_module():
    mp = pytest.MonkeyPatch()
    yield mp
    mp.undo()


def csrf(c, path):
    return re.search(r'name="csrf-token" content="([^"]+)"', c.get(path).text).group(1)


def test_full_flow(client, fake_services):
    c = client
    r = c.get("/login", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/register"  # first run: set up the administrator
    token = csrf(c, "/register")
    r = c.post("/register", data={"username": "alice", "password": "correct horse battery", "csrf": token},
               follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/settings"

    token = csrf(c, "/settings")
    r = c.post("/settings/token", data={"token": "wrong", "token_type": "access_token", "csrf": token})
    assert "rejected the token" in r.text
    r = c.post("/settings/token", data={"token": fake_services.TOKEN, "token_type": "access_token", "csrf": token})
    assert "Connected to GitLab as alice" in r.text and fake_services.TOKEN not in r.text

    h = {"X-CSRF-Token": csrf(c, "/")}
    assert c.get("/api/health").json()["llm"]["ok"]
    projects = c.get("/api/projects").json()
    assert projects[0]["path"] == "grp/app"
    branch = c.get("/api/projects/1/branches").json()[0]["name"]
    commit = c.get(f"/api/projects/1/commits?ref={branch}").json()[0]
    scan = c.get(f"/api/projects/1/commits/{commit['sha']}/scan?path=grp/app").json()
    assert scan["state"] == "finished" and scan["finding_count"] == 2

    assert c.post("/api/jobs", json={"project_id": 1, "project_path": "grp/app", "branch": branch,
                                     "sha": commit["sha"]}).status_code == 403  # no CSRF header
    job_id = c.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": branch,
                                                  "sha": commit["sha"], "severities": ["high", "low"]}).json()["id"]

    from app.analyzer import JobRunner
    from app.worker import _claim_next
    assert _claim_next() == job_id
    JobRunner(job_id).run()

    job = c.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "done", job["error"]
    real, fp = job["findings"]  # sorted by severity
    assert real["uuid"] == "uuid-1" and real["verdict"] == "true_positive"
    assert real["code_flows"][0] == {"file": "app.py", "line": 12, "kind": "source"}
    assert "+    cur.execute(\"SELECT * FROM users WHERE name = ?\", (name,))" in real["diff"]
    assert real["validation_status"] == "not_configured"
    assert real["injection_suspected"] is True
    assert fp["verdict"] == "false_positive" and fp["fix"] is None

    r = c.post(f"/api/findings/{real['id']}/mr", headers=h).json()
    assert r["iid"] == 7
    committed = fake_services.gitlab.state["last_commit"]
    assert committed["actions"][0]["file_path"] == "app.py"
    assert "(name,)" in committed["actions"][0]["content"]
    # One merge request per finding: the button turns into "opened" and a second click is refused.
    fake_services.gitlab.state["last_commit"] = None
    r = c.post(f"/api/findings/{real['id']}/mr", headers=h)
    assert r.status_code == 409 and "merge request !7 was already opened by alice" in r.json()["detail"]
    assert fake_services.gitlab.state["last_commit"] is None
    real = c.get(f"/api/jobs/{job_id}").json()["findings"][0]
    assert real["mr"]["mr_iid"] == 7 and real["mr"]["url"].endswith("/merge_requests/7")

    # First post creates a confidential issue assigned to the poster, then comments on it.
    r = c.post(f"/api/findings/{fp['id']}/note", headers=h).json()
    assert r["created"] is True and r["where"] == "new issue #1 (assigned to you)"
    issue = fake_services.gitlab.state["issues"][0]
    assert issue["assignee_ids"] == [1] and issue["confidential"] is True
    assert "[Low] Hard-coded value in query in app.py:17" == issue["title"]
    assert f"vulntriage-finding: {fp['uuid']}" in issue["description"] and len(issue["notes"]) == 1
    # The same analysis can't be posted twice; the job shows where it went.
    r = c.post(f"/api/findings/{fp['id']}/note", headers=h)
    assert r.status_code == 409 and "already posted on issue #1 by alice" in r.json()["detail"]
    assert len(issue["notes"]) == 1
    fp = c.get(f"/api/jobs/{job_id}").json()["findings"][1]
    assert fp["posted"]["issue_iid"] == 1 and fp["posted"]["url"].endswith("#note_501")
    from app.config import settings
    assert fp["gitlab_url"] == f"{settings.gitlab_url}/grp/app/-/pipelines/10/security"
    # A new analysis (run again) finds that issue and only adds a comment.
    rerun_id = c.post(f"/api/jobs/{job_id}/rerun", headers=h).json()["id"]
    assert _claim_next() == rerun_id
    JobRunner(rerun_id).run()
    fp2 = c.get(f"/api/jobs/{rerun_id}").json()["findings"][1]
    assert fp2["copied_from"]["job"] == job_id and fp2["verdict"] == fp["verdict"]  # same code: not asked again
    # The copied analysis is the one already posted, so it isn't posted again.
    assert fp2["uuid"] == fp["uuid"] and fp2["posted"]["issue_iid"] == 1 and fp2["posted"]["copied_from_job"] == job_id
    assert fp2["posted"]["url"] == fp["posted"]["url"]
    r = c.post(f"/api/findings/{fp2['id']}/note", headers=h)
    assert r.status_code == 409 and f"from analysis #{job_id}" in r.json()["detail"] and len(issue["notes"]) == 1
    # The same goes for the fix merge request of a copied analysis.
    real2 = c.get(f"/api/jobs/{rerun_id}").json()["findings"][0]
    assert real2["copied_from"]["job"] == job_id and real2["mr"]["mr_iid"] == 7
    assert real2["mr"]["copied_from_job"] == job_id
    r = c.post(f"/api/findings/{real2['id']}/mr", headers=h)
    assert r.status_code == 409 and f"!7 was already opened from analysis #{job_id}" in r.json()["detail"]
    assert fake_services.gitlab.state["last_commit"] is None
    # A new result from the model can be posted: it finds that issue and only adds a comment.
    from app.db import Finding, SessionLocal
    with SessionLocal() as db:
        db.get(Finding, fp2["id"]).copied_from = None
        db.commit()
    assert c.get(f"/api/jobs/{rerun_id}").json()["findings"][1]["posted"] is None
    r = c.post(f"/api/findings/{fp2['id']}/note", headers=h).json()
    assert r["created"] is False and r["issue_iid"] == 1 and r["url"].endswith("#note_502")
    assert len(fake_services.gitlab.state["issues"]) == 1 and len(issue["notes"]) == 2
    r = c.post(f"/api/findings/{fp['id']}/dismiss", headers=h)
    assert r.status_code == 409  # no vulnerability record on a feature branch

    _check_report(c, job_id, fake_services)

    page = c.get(f"/jobs/{job_id}")
    assert page.status_code == 200 and "Analysis #" in page.text
    assert 'id="post-selected"' in page.text and 'id="pick-all"' in page.text and "GitLab issue" in page.text


def test_validation_runs_configured_commands(client, fake_services):
    from sqlalchemy import update
    from app.db import Finding, ProjectConfig, SessionLocal
    from app.analyzer import JobRunner
    from app.worker import _claim_next
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key="", code_key="-"))  # force a fresh analysis
        db.add(ProjectConfig(gitlab_project_id=1, test_cmd="python3 -c \"import ast,sys; ast.parse(open('app.py').read())\"",
                             rescan_cmd="! grep -n \"'%s'\" {files}"))
        db.commit()
    h = {"X-CSRF-Token": csrf(client, "/")}
    sha = fake_services.SHA
    job_id = client.post(f"/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app",
                                                        "branch": "feature/login", "sha": sha}).json()["id"]
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    real = client.get(f"/api/jobs/{job_id}").json()["findings"][0]
    assert real["verdict"] == "true_positive"
    assert real["validation_status"] == "passed", real["validation_log"]
    assert "[rescan: ok]" in real["validation_log"]


def test_cache_reuses_identical_analysis(client, fake_services):
    from app.analyzer import JobRunner
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app",
                                                       "branch": "feature/login", "sha": fake_services.SHA}).json()["id"]
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    real = client.get(f"/api/jobs/{job_id}").json()["findings"][0]
    assert real["validation_status"] == "passed"


def test_default_severities_skip_low(client, fake_services):
    from app.analyzer import JobRunner
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app",
                                                       "branch": "feature/login", "sha": fake_services.SHA}).json()["id"]
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["severities"] == ["critical", "high"]
    high, low = job["findings"]
    assert high["severity"] == "high" and high["status"] == "done"
    assert low["severity"] == "low" and low["status"] == "skipped" and "wasn't selected" in low["reasoning"]
    assert not low["votes"]


def test_bad_severities_rejected(client, fake_services):
    h = {"X-CSRF-Token": csrf(client, "/")}
    for bad in ([], ["bogus"]):
        r = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app",
                                                      "branch": "feature/login", "sha": fake_services.SHA,
                                                      "severities": bad})
        assert r.status_code == 400


def test_resume_keeps_finished_findings(client, fake_services, monkeypatch):
    """A job interrupted mid-analysis (restart, upgrade) keeps its results and continues where it stopped."""
    from sqlalchemy import update
    from app.analyzer import JobRunner
    from app.db import AnalysisJob, Finding, SessionLocal
    from app.llm_client import LLMClient
    from app.worker import _claim_next, recover_interrupted
    h = {"X-CSRF-Token": csrf(client, "/")}
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key="", code_key="-"))  # no reuse of earlier results
        db.commit()
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                                       "sha": fake_services.SHA, "severities": ["high", "low"]}).json()["id"]
    assert _claim_next() == job_id

    # Let the High finding finish and the Low one get one triage vote, then "kill the container".
    # The findings and votes run in parallel, so count the Low finding's votes under a lock.
    import threading
    real_chat, calls, lock = LLMClient.chat_json, [], threading.Lock()
    crash = {"on": True, "low_votes": 0}

    def chat(self, system, user, *a, **kw):
        low_vote = "app.py:17" in user and "triaging" in system
        with lock:
            calls.append("triage" if "triaging" in system else "fix")
            crash["low_votes"] += low_vote
            stop = crash["on"] and low_vote and crash["low_votes"] > 1
        if stop:
            raise KeyboardInterrupt("container stopped")
        return real_chat(self, system, user, *a, **kw)
    monkeypatch.setattr(LLMClient, "chat_json", chat)
    with pytest.raises(KeyboardInterrupt):
        JobRunner(job_id).run()
    with SessionLocal() as db:
        high, low = db.get(AnalysisJob, job_id).findings
        assert high.status == "done" and high.verdict == "true_positive"
        assert low.status == "analyzing" and len(low.llm_log) == 1  # one vote saved before the stop
        high_id = high.id

    recover_interrupted()  # what the worker does when it starts again
    calls.clear()
    crash["on"] = False
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "done"
    high, low = job["findings"]
    assert high["id"] == high_id and high["verdict"] == "true_positive"  # kept, not redone
    assert low["verdict"] == "false_positive"
    assert calls == ["triage", "triage"]  # only the two missing votes of the interrupted finding


def test_admin_llm_settings(client):
    token = csrf(client, "/admin")
    assert "Local model (vLLM)" not in client.get("/settings").text  # lives on the Admin page now
    page = client.get("/admin").text
    assert "Local model (vLLM)" in page
    base = re.search(r'name="base_url" value="([^"]+)"', page).group(1)
    r = client.post("/admin/llm", data={"base_url": base, "model": "other-model", "api_key": "sekrit-key-123",
                                        "structured_output": "1", "timeout_s": "120", "csrf": token})
    assert "Model settings saved" in r.text and "sekrit-key-123" not in r.text
    health = client.get("/api/health").json()["llm"]
    assert not health["ok"] and "other-model" in health["detail"]
    from app.llm_client import llm_config
    cfg = llm_config()
    assert cfg["api_key"] == "sekrit-key-123" and cfg["timeout_s"] == 120
    assert cfg["parallel"] == 4  # left blank: the .env default
    r = client.post("/admin/llm", data={"base_url": base, "model": "other-model", "timeout_s": "120", "parallel": "2",
                                        "csrf": token})
    assert "Model settings saved" in r.text and 'name="parallel" type="number" min="1" max="32" value="2"' in r.text
    assert llm_config()["parallel"] == 2
    assert llm_config()["max_tokens"] == 4096
    r = client.post("/admin/llm", data={"base_url": base, "model": "other-model", "timeout_s": "120", "max_tokens": "8192",
                                        "csrf": token})
    assert 'name="max_tokens" type="number" min="1024" max="32768" value="8192"' in r.text
    assert llm_config()["max_tokens"] == 8192
    for bad in ("512", "40000", "x"):
        r = client.post("/admin/llm", data={"base_url": base, "model": "m", "timeout_s": "120", "max_tokens": bad,
                                            "csrf": token})
        assert "reply limit between 1024 and 32768 tokens" in r.text
    assert llm_config()["context_budget"] == 100000
    r = client.post("/admin/llm", data={"base_url": base, "model": "other-model", "timeout_s": "120",
                                        "context_budget": "60,000", "csrf": token})
    assert 'name="context_budget" type="number" min="4000" max="1000000" step="1000" value="60000"' in r.text
    assert llm_config()["context_budget"] == 60000
    for bad in ("1000", "2000000", "x"):
        r = client.post("/admin/llm", data={"base_url": base, "model": "m", "timeout_s": "120", "context_budget": bad,
                                            "csrf": token})
        assert "code budget between 4000 and 1000000 tokens" in r.text
    assert llm_config()["context_budget"] == 60000
    for bad in ("0", "33", "x"):
        r = client.post("/admin/llm", data={"base_url": base, "model": "m", "timeout_s": "120", "parallel": bad,
                                            "csrf": token})
        assert "between 1 and 32 parallel requests" in r.text
    r = client.post("/admin/llm", data={"base_url": "ftp://nope", "model": "x", "timeout_s": "120", "csrf": token})
    assert "must start with http" in r.text
    r = client.post("/admin/llm/reset", data={"csrf": token})
    assert "reset to the values in .env" in r.text
    assert client.get("/api/health").json()["llm"]["ok"]
    assert llm_config()["context_budget"] == 100000


@pytest.fixture
def directory(monkeypatch):
    """ldap3's in-memory directory with a service account and two Windows users."""
    from ldap3 import MOCK_SYNC, Connection, Server
    from app import ldap_auth
    server = Server("mock")
    seed = Connection(server, client_strategy=MOCK_SYNC)
    seed.strategy.add_entry("CN=svc,DC=corp,DC=test", {"userPassword": "svc-pass", "objectClass": "person"})
    for sam, name in (("jdoe", "John Doe"), ("alice", "Alice Local")):
        seed.strategy.add_entry(f"CN={name},OU=Staff,DC=corp,DC=test", {
            "userPassword": f"{sam}-windows-pw", "objectClass": ["top", "person", "user"], "objectCategory": "person",
            "sAMAccountName": sam, "displayName": name, "mail": f"{sam}@corp.test"})
    monkeypatch.setattr(ldap_auth, "_server", lambda cfg: server)
    monkeypatch.setattr(ldap_auth, "_CONNECTION_OPTIONS", {"client_strategy": MOCK_SYNC})
    return server


def test_windows_sign_in(client, directory):
    from fastapi.testclient import TestClient
    from app.main import app
    token = csrf(client, "/admin")
    settings_form = {"enabled": "1", "server_url": "ldaps://dc01.corp.test:636", "domain": "corp.test",
                     "base_dn": "DC=corp,DC=test", "bind_dn": "CN=svc,DC=corp,DC=test", "bind_password": "svc-pass",
                     "user_filter": "", "verify_cert": "1", "csrf": token}
    r = client.post("/admin/ldap", data={**settings_form, "server_url": "http://nope"})
    assert "server URL must look like" in r.text
    r = client.post("/admin/ldap", data=settings_form)
    assert "Windows sign-in is on" in r.text and "svc-pass" not in r.text
    r = client.post("/admin/ldap/test", data={"lookup": "CORP\\jdoe", "csrf": token})
    assert "Test passed" in r.text and "CN=John Doe" in r.text

    with TestClient(app) as w:
        page = w.get("/login").text
        assert "Windows account" in page and "/register" not in page  # no self-service sign-up
        token = csrf(w, "/login")
        assert "wrong" in w.post("/login", data={"username": "jdoe", "password": "nope", "csrf": token}).text
        r = w.post("/login", data={"username": "jdoe", "password": "", "csrf": token})  # never an anonymous bind
        assert r.status_code == 422 or "wrong" in r.text
        from app import ldap_auth
        from app.db import SessionLocal
        with SessionLocal() as db, pytest.raises(ldap_auth.LDAPAuthError):
            ldap_auth.authenticate(ldap_auth.ldap_config(db), "jdoe", "")
        r = w.post("/login", data={"username": "CORP\\jdoe", "password": "jdoe-windows-pw", "csrf": token},
                   follow_redirects=False)
        assert r.status_code == 303  # first sign-in creates the profile
        settings_page = w.get("/settings").text
        assert "Windows account" in settings_page and 'action="/settings/password"' not in settings_page
        token = csrf(w, "/settings")
        r = w.post("/settings/password", data={"current": "x", "new": "y" * 12, "csrf": token})
        assert "change its password in Windows" in r.text
        assert w.get("/admin").status_code == 403

        # A local account with the same name is never taken over by Windows sign-in.
        token = csrf(w, "/login")
        w.post("/logout", data={"csrf": csrf(w, "/settings")})
        r = w.post("/login", data={"username": "alice", "password": "alice-windows-pw", "csrf": csrf(w, "/login")})
        assert "wrong" in r.text

    from app.db import SessionLocal, User
    from sqlalchemy import select
    with SessionLocal() as db:
        u = db.scalar(select(User).where(User.username == "jdoe"))
        assert u.is_windows and u.display_name == "John Doe" and u.email == "jdoe@corp.test" and not u.is_admin
        assert u.password_hash == "windows"
    r = client.post("/admin/ldap", data={**settings_form, "enabled": "", "bind_password": "",
                                         "csrf": csrf(client, "/admin")})
    assert "Windows sign-in is off" in r.text
    from app import ldap_auth
    with SessionLocal() as db:
        assert ldap_auth.ldap_config(db)["bind_password"] == "svc-pass"  # blank kept the saved password


def test_registration_closed_after_first_admin(client):
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as anon:
        assert anon.get("/register", follow_redirects=False).headers["location"] == "/login"
        token = csrf(anon, "/login")
        r = anon.post("/register", data={"username": "mallory", "password": "x" * 14, "csrf": token})
        assert r.status_code == 403


def test_admin_pages_and_user_creation(client):
    token = csrf(client, "/admin")
    r = client.post("/admin/users", data={"username": "bob", "password": "another long password", "csrf": token})
    assert r.status_code == 200 and "Added bob" in r.text
    r = client.post("/admin/users", data={"username": "bob", "password": "another long password", "csrf": token})
    assert "is taken" in r.text
    assert client.get("/history").status_code == 200
    assert "System status" in client.get("/settings").text and client.get("/").status_code == 200


def test_non_admin_cannot_change_llm_settings(fake_services, client):
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as bob:
        token = csrf(bob, "/login")
        bob.post("/login", data={"username": "bob", "password": "another long password", "csrf": token})
        token = csrf(bob, "/settings")
        assert "Local model (vLLM)" not in bob.get("/settings").text
        assert bob.get("/admin").status_code == 403
        for path, data in (("/admin/llm", {"base_url": "http://evil:1/v1", "model": "m", "timeout_s": "60"}),
                           ("/admin/llm/reset", {}), ("/admin/ldap", {"enabled": "1"}), ("/admin/ldap/test", {})):
            assert bob.post(path, data={**data, "csrf": token}).status_code == 403
        assert "models" not in bob.get("/api/health").json()["llm"]  # served models are for admins
        assert bob.get("/jobs/1/report.xlsx").status_code == 400  # no GitLab token: can't prove project access


def _check_report(c, job_id, fake_services):
    """The Excel report lists every finding with its posted state and links back to GitLab."""
    from io import BytesIO
    from openpyxl import load_workbook
    r = c.get(f"/jobs/{job_id}/report.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f'filename="securefix-grp-app-{fake_services.SHA[:8]}-analysis-{job_id}.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Findings"]
    ws = wb["Findings"]
    head = [cell.value for cell in ws[1]]
    assert head[:10] == ["#", "Severity", "Verdict", "Confidence", "Finding", "CWE", "Location", "GitLab issue",
                         "Link", "Analysis posted"]
    job = c.get(f"/api/jobs/{job_id}").json()
    rows = {row[4].value: row for row in ws.iter_rows(min_row=2)}
    assert len(rows) == len(job["findings"])
    col = {name: i for i, name in enumerate(head)}
    for f in job["findings"]:
        row = rows[f["title"]]
        assert row[col["Severity"]].value == f["severity"].capitalize()
        assert row[col["Link"]].hyperlink.target == f["gitlab_url"]
        assert row[col["Location"]].value.startswith(f"{f['file']}:{f['start_line']}")
        if f["posted"]:
            assert row[col["Analysis posted"]].value == "Yes"
            assert row[col["GitLab issue"]].value == f"#{f['posted']['issue_iid']}"
            assert row[col["GitLab issue"]].hyperlink.target == f["posted"]["url"].split("#note_")[0]
        else:
            assert row[col["Analysis posted"]].value == "No" and not row[col["GitLab issue"]].value
    real = next(f for f in job["findings"] if f["verdict"] == "true_positive")
    assert rows[real["title"]][col["Fix MR"]].value == "!7"
    assert "SecureFix analysis report" == wb["Summary"]["B2"].value


def test_report_text_is_never_a_formula():
    from io import BytesIO
    from types import SimpleNamespace as NS
    from openpyxl import load_workbook
    from app.report import build_report
    f = NS(id=1, uuid="u", vulnerability_gid=None, title="=HYPERLINK(\"http://evil\",\"x\")", severity="high",
           scanner="semgrep", cwe="CWE-89", file="app.py", start_line=3, end_line=3, status="done",
           verdict="true_positive", confidence=0.9, reasoning="+cmd|' /C calc'!A0", fix=None, diff="",
           validation_status=None, error=None, actions=[], copied_from=None, duration_s=None, completion_tokens=None)
    job = NS(id=5, project_path="grp/app", branch="main", sha="a" * 40, pipeline_id=None, status="done",
             user=NS(username="alice"), created_at=None, finished_at=None, severities=None, thinking=None, budget=100000,
             findings=[f])
    ws = load_workbook(BytesIO(build_report(job, "https://gl", "alice")))["Findings"]
    assert ws["E2"].data_type == "s" and ws["E2"].value.startswith("=HYPERLINK")
    assert ws["V2"].data_type == "s" and ws["V2"].value == "+cmd|' /C calc'!A0"  # control character dropped


def test_vulnerability_url():
    from types import SimpleNamespace as NS
    from app.prompts import vulnerability_url
    job = NS(project_path="grp/app", pipeline_id=10)
    assert vulnerability_url(NS(vulnerability_gid="gid://gitlab/Vulnerability/42"), job, "https://gl") \
        == "https://gl/grp/app/-/security/vulnerabilities/42"
    assert vulnerability_url(NS(vulnerability_gid=None), job, "https://gl") == "https://gl/grp/app/-/pipelines/10/security"
    assert vulnerability_url(NS(vulnerability_gid=None), NS(project_path="grp/app", pipeline_id=None), "https://gl") is None


def _run_fresh_job(client, h, fake_services, parallel):
    """Runs the High + Low findings of the fixture commit from scratch with the given parallel setting."""
    from sqlalchemy import update
    from app import app_settings
    from app.analyzer import JobRunner
    from app.db import Finding, SessionLocal
    from app.worker import _claim_next
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key="", code_key="-"))  # no reuse of earlier results
        app_settings.save(db, "llm", {"parallel": str(parallel)})
        db.commit()
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                                       "sha": fake_services.SHA, "severities": ["high", "low"]}).json()["id"]
    assert _claim_next() == job_id
    load = fake_services.vllm.state.load
    load.update(peak=0, delay=0.2)
    try:
        JobRunner(job_id).run()
    finally:
        load["delay"] = 0.0
    return job_id, load["peak"]


def test_parallel_requests_same_results(client, fake_services):
    from app import app_settings
    from app.db import SessionLocal
    h = {"X-CSRF-Token": csrf(client, "/")}
    try:
        serial_id, serial_peak = _run_fresh_job(client, h, fake_services, 1)
        parallel_id, parallel_peak = _run_fresh_job(client, h, fake_services, 4)
    finally:
        with SessionLocal() as db:
            app_settings.clear(db, "llm")
    assert serial_peak == 1
    assert 1 < parallel_peak <= 4
    keys = ("uuid", "status", "verdict", "confidence", "votes", "reasoning", "evidence", "diff", "validation_status")
    serial, parallel = (client.get(f"/api/jobs/{i}").json() for i in (serial_id, parallel_id))
    assert parallel["status"] == "done"
    assert [{k: f[k] for k in keys} for f in serial["findings"]] == [{k: f[k] for k in keys} for f in parallel["findings"]]


def test_cancel_while_parallel(client, fake_services, monkeypatch):
    """Cancelling stops the parallel requests; unfinished findings stay pending, nothing is marked as an error."""
    from sqlalchemy import update
    from app import app_settings
    from app.analyzer import JobRunner
    from app.db import AnalysisJob, Finding, SessionLocal
    from app.llm_client import LLMClient
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key="", code_key="-"))
        app_settings.save(db, "llm", {"parallel": "4"})
        db.commit()
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                                       "sha": fake_services.SHA, "severities": ["high", "low"]}).json()["id"]
    assert _claim_next() == job_id
    real_chat = LLMClient.chat_json

    def chat(self, *a, **kw):
        with SessionLocal() as db:  # the user presses Cancel while the first requests are out
            db.execute(update(AnalysisJob).where(AnalysisJob.id == job_id).values(status="cancelled"))
            db.commit()
        return real_chat(self, *a, **kw)
    monkeypatch.setattr(LLMClient, "chat_json", chat)
    try:
        JobRunner(job_id).run()
    finally:
        with SessionLocal() as db:
            app_settings.clear(db, "llm")
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "cancelled" and job["progress"] == "Cancelled"
    statuses = {f["status"] for f in job["findings"]}
    assert statuses <= {"pending", "done"} and "pending" in statuses, statuses


def _forget_results():
    """Earlier tests' analyses can't be reused or copied."""
    from sqlalchemy import update
    from app.db import Finding, SessionLocal
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key="", code_key="-"))
        db.commit()


def _new_job(client, h, fake_services):
    from app.worker import _claim_next
    job_id = client.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                                       "sha": fake_services.SHA, "severities": ["high", "low"]}).json()["id"]
    assert _claim_next() == job_id
    return job_id


def _count_calls(monkeypatch, fail=None):
    """Counts model calls by finding line; fail(system, user, n) may raise instead of answering."""
    import threading
    from app.llm_client import LLMClient
    real, calls, lock = LLMClient.chat_json, [], threading.Lock()

    def chat(self, system, user, *a, **kw):
        with lock:
            calls.append(("triage" if "triaging" in system else "fix", "app.py:17" in user and "low" or "high"))
            n = len([c for c in calls if c == calls[-1]])
        if fail:
            fail(system, user, n)
        return real(self, system, user, *a, **kw)
    monkeypatch.setattr(LLMClient, "chat_json", chat)
    return calls


def test_copies_analysis_from_earlier_scan(client, fake_services, monkeypatch):
    """A later analysis with the same vulnerability, lines and code copies the earlier result instead of asking the model."""
    from sqlalchemy import update
    from app.analyzer import JobRunner
    from app.db import AnalysisJob, Finding, SessionLocal
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    calls = _count_calls(monkeypatch)
    first = _new_job(client, h, fake_services)
    JobRunner(first).run()
    assert calls and client.get(f"/api/jobs/{first}").json()["findings"][0]["copied_from"] is None
    with SessionLocal() as db:  # as if analyzed before this version: no code key and no exact cache match
        db.execute(update(Finding).where(Finding.job_id == first).values(code_key="", cache_key=""))
        db.commit()

    calls.clear()
    second = _new_job(client, h, fake_services)
    JobRunner(second).run()
    assert calls == []
    old, new = (client.get(f"/api/jobs/{j}").json()["findings"] for j in (first, second))
    for o, n in zip(old, new):
        assert n["status"] == "done" and n["copied_from"]["job"] == first and n["copied_from"]["finding"] == o["id"]
        assert n["copied_from"]["sha"] == fake_services.SHA
        assert (n["verdict"], n["confidence"], n["reasoning"], n["diff"]) == (o["verdict"], o["confidence"], o["reasoning"], o["diff"])
        assert n["posted"] is None  # the earlier analysis wasn't posted either
    with SessionLocal() as db:  # the old findings' keys were worked out from that commit's files
        keys = {f.code_key for f in db.get(AnalysisJob, first).findings}
        assert "" not in keys and "-" not in keys
    page = client.get(f"/jobs/{second}")
    assert page.status_code == 200

    # Posting the earlier analysis after it was copied: the copy shows it as posted and refuses a second comment.
    r = client.post(f"/api/findings/{old[0]['id']}/note", headers=h)
    assert r.status_code == 200
    notes = len(fake_services.gitlab.state["issues"][-1]["notes"])
    new = client.get(f"/api/jobs/{second}").json()["findings"]
    assert new[0]["posted"]["url"] == r.json()["url"] and new[0]["posted"]["copied_from_job"] == first
    assert all(n["posted"] is None for n in new[1:])
    r = client.post(f"/api/findings/{new[0]['id']}/note", headers=h)
    assert r.status_code == 409 and len(fake_services.gitlab.state["issues"][-1]["notes"]) == notes

    # Code changed (different key) or the finding moved: the model analyzes it again.
    with SessionLocal() as db:
        db.execute(update(Finding).values(code_key="changed", cache_key=""))
        db.commit()
    third = _new_job(client, h, fake_services)
    JobRunner(third).run()
    assert ("triage", "high") in calls and ("triage", "low") in calls
    assert all(f["copied_from"] is None for f in client.get(f"/api/jobs/{third}").json()["findings"])


def test_cut_off_reply_retried_with_more_tokens(client, fake_services):
    from app import prompts
    from app.llm_client import SHORTER, BadReply, LLMClient
    cut = fake_services.vllm.state.cut
    llm = LLMClient()
    try:
        cut.update(left=1, seen=[])
        res, _ = llm.chat_json(prompts.TRIAGE_SYSTEM, "Location: app.py:7-7", prompts.TRIAGE_SCHEMA, "triage")
        assert res["verdict"] == "true_positive"
        assert [t for t, _ in cut["seen"]] == [4096, 8192] and cut["seen"][1][1].endswith(SHORTER)
        cut.update(left=2, seen=[])
        with pytest.raises(BadReply, match="cut off at 8192 tokens"):
            llm.chat_json(prompts.TRIAGE_SYSTEM, "Location: app.py:7-7", prompts.TRIAGE_SCHEMA, "triage")
    finally:
        cut.update(left=0, seen=[])
        llm.close()


def test_one_unreadable_vote_does_not_fail_the_finding(client, fake_services, monkeypatch):
    from app.analyzer import JobRunner
    from app.llm_client import BadReply
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()

    def fail(system, user, n):
        if "triaging" in system and "app.py:17" in user and n == 1:
            raise BadReply("Model reply wasn't valid JSON: Unterminated string starting at: line 7 column 16")
    _count_calls(monkeypatch, fail)
    job_id = _new_job(client, h, fake_services)
    JobRunner(job_id).run()
    low = client.get(f"/api/jobs/{job_id}").json()["findings"][1]
    assert low["status"] == "done" and low["verdict"] == "false_positive"
    assert len(low["votes"]) == 2 and low["confidence"] == round(0.8 * 2 / 3, 3)
    assert low["reasoning"].startswith("1 of the model's 3 triage runs returned an unreadable reply")


def test_retry_failed_findings(client, fake_services, monkeypatch):
    """Only the failed findings run again, in the same analysis; finished ones are kept as they are."""
    from app.analyzer import JobRunner
    from app.db import AnalysisJob, SessionLocal
    from app.llm_client import LLMError
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    broken = {"on": True}

    def fail(system, user, n):
        if broken["on"] and "app.py:17" in user:
            raise LLMError("vLLM returned 500: out of memory")
    calls = _count_calls(monkeypatch, fail)
    job_id = _new_job(client, h, fake_services)
    JobRunner(job_id).run()
    high, low = client.get(f"/api/jobs/{job_id}").json()["findings"]
    assert high["status"] == "done" and low["status"] == "error" and "out of memory" in low["error"]
    assert 'id="retry-failed"' in client.get(f"/jobs/{job_id}").text

    with SessionLocal() as db:  # not while it runs
        db.get(AnalysisJob, job_id).status = "running"
        db.commit()
    assert client.post(f"/api/jobs/{job_id}/retry-failed", headers=h).status_code == 409
    with SessionLocal() as db:
        db.get(AnalysisJob, job_id).status = "done"
        db.commit()
    assert client.post(f"/api/jobs/{job_id}/retry-failed").status_code == 403  # CSRF
    r = client.post(f"/api/jobs/{job_id}/retry-failed", headers=h).json()
    assert r == {"status": "queued", "retried": 1}

    broken["on"] = False
    calls.clear()
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "done"
    high2, low2 = job["findings"]
    assert high2 == high  # untouched
    assert low2["status"] == "done" and low2["verdict"] == "false_positive" and low2["error"] is None
    assert calls == [("triage", "low")] * 3
    assert client.post(f"/api/jobs/{job_id}/retry-failed", headers=h).status_code == 409  # nothing left to retry


def _split_low_vote(monkeypatch):
    """The first triage run of the low finding says it is real, so its runs disagree; returns the model calls."""
    import threading
    from app.llm_client import LLMClient
    real, calls, lock, seen = LLMClient.chat_json, [], threading.Lock(), {"low": 0}

    def chat(self, system, user, *a, **kw):
        res, raw = real(self, system, user, *a, **kw)
        with lock:
            calls.append(("triage" if "triaging" in system else "fix", "app.py:17" in user and "low" or "high"))
            if calls[-1] == ("triage", "low"):
                seen["low"] += 1
                if seen["low"] == 1:
                    res = {**res, "verdict": "true_positive", "confidence": 0.6,
                           "reasoning": "The value could come from the request."}
        return res, raw
    monkeypatch.setattr(LLMClient, "chat_json", chat)
    return calls


def test_needs_review_confirmed_as_real(client, fake_services, monkeypatch):
    """A reviewer sees every run's argument, confirms the finding as real, and only its fix is written."""
    from app.analyzer import JobRunner
    from app.db import AnalysisJob, SessionLocal
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    calls = _split_low_vote(monkeypatch)
    job_id = _new_job(client, h, fake_services)
    JobRunner(job_id).run()
    high, low = client.get(f"/api/jobs/{job_id}").json()["findings"]
    assert low["verdict"] == "needs_review" and low["fix"] is None and low["reviewed"] is None
    assert sorted(v["verdict"] for v in low["votes"]) == ["false_positive", "false_positive", "true_positive"]
    assert all(v["reasoning"] and v["source"] and v["sink"] for v in low["votes"])

    assert client.post(f"/api/findings/{low['id']}/confirm").status_code == 403  # CSRF
    assert client.post(f"/api/findings/{high['id']}/confirm", headers=h).status_code == 409  # already real
    with SessionLocal() as db:  # not while it runs
        db.get(AnalysisJob, job_id).status = "running"
        db.commit()
    assert client.post(f"/api/findings/{low['id']}/confirm", headers=h).status_code == 409
    with SessionLocal() as db:
        db.get(AnalysisJob, job_id).status = "done"
        db.commit()
    assert client.post(f"/api/findings/{low['id']}/confirm", headers=h).json() == {"status": "queued"}
    queued = client.get(f"/api/jobs/{job_id}").json()["findings"][1]
    assert queued["status"] == "pending" and queued["verdict"] == "true_positive"
    assert queued["reviewed"]["by"] == "alice" and queued["reviewed"]["from_verdict"] == "needs_review"

    calls.clear()
    assert _claim_next() == job_id
    JobRunner(job_id).run()
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "done"
    high2, low2 = job["findings"]
    assert high2 == high
    assert calls and set(calls) == {("fix", "low")}  # no triage again (a failed validation retries the fix)
    assert low2["status"] == "done" and low2["verdict"] == "true_positive" and low2["error"] is None
    assert "(name,)" in low2["diff"] and low2["validation_status"]
    assert low2["votes"] == low["votes"] and low2["reasoning"] == low["reasoning"]
    assert client.post(f"/api/findings/{low['id']}/confirm", headers=h).status_code == 409  # settled

    iid = client.post(f"/api/findings/{low['id']}/note", headers=h).json()["issue_iid"]
    note = next(i for i in fake_services.gitlab.state["issues"] if i["iid"] == iid)["notes"][-1]
    assert note.startswith("**Real vulnerability: confirmed by alice after the automated triage needed review**")


def test_needs_review_marked_false_positive(client, fake_services, monkeypatch):
    """Marking a Needs review finding as a false positive dismisses it in GitLab, and a later scan copies the decision."""
    from app.analyzer import JobRunner
    from app.db import Finding, SessionLocal
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    _split_low_vote(monkeypatch)
    job_id = _new_job(client, h, fake_services)
    JobRunner(job_id).run()
    high, low = client.get(f"/api/jobs/{job_id}").json()["findings"]
    assert low["verdict"] == "needs_review"
    assert client.post(f"/api/findings/{high['id']}/false-positive", headers=h).status_code == 409
    gid = "gid://gitlab/Vulnerability/42"
    with SessionLocal() as db:
        db.get(Finding, low["id"]).vulnerability_gid = gid
        db.commit()
    fake_services.gitlab.state["dismissed"] = []
    r = client.post(f"/api/findings/{low['id']}/false-positive", headers=h).json()
    assert r == {"ok": True, "dismissed": True}
    assert [d["id"] for d in fake_services.gitlab.state["dismissed"]] == [gid]
    low = client.get(f"/api/jobs/{job_id}").json()["findings"][1]
    assert low["verdict"] == "false_positive" and low["reviewed"]["from_verdict"] == "needs_review"
    assert low["actions"][-1]["type"] == "marked as false positive and dismissed in GitLab"

    later = _new_job(client, h, fake_services)  # same code: copies the reviewed verdict
    JobRunner(later).run()
    copy = client.get(f"/api/jobs/{later}").json()["findings"][1]
    assert copy["copied_from"]["finding"] == low["id"]
    assert copy["verdict"] == "false_positive" and copy["reviewed"]["by"] == "alice"


def test_shared_analysis_is_read_only(client, fake_services):
    """Other users who are Developer or higher on the project can view an analysis and export it, but not act on it."""
    from app import main
    from app.analyzer import JobRunner
    from app.main import app
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    job_id = _new_job(client, h, fake_services)
    JobRunner(job_id).run()
    findings = client.get(f"/api/jobs/{job_id}").json()["findings"]
    real = next(f for f in findings if f["verdict"] == "true_positive")
    token = csrf(client, "/admin")
    client.post("/admin/users", data={"username": "dave", "password": "dave's long password", "csrf": token})

    with TestClient(app) as dave:
        dave.post("/login", data={"username": "dave", "password": "dave's long password", "csrf": csrf(dave, "/login")})
        assert "Add your GitLab token" in dave.get("/history?scope=shared").text
        dave.post("/settings/token", data={"token": fake_services.BOB_TOKEN, "token_type": "access_token",
                                           "csrf": csrf(dave, "/settings")})
        dh = {"X-CSRF-Token": csrf(dave, "/")}
        assert f"/jobs/{job_id}\"" not in dave.get("/history").text  # Mine: only dave's own
        shared = dave.get("/history?scope=shared").text
        assert f"/jobs/{job_id}\"" in shared and "alice" in shared
        everyone = dave.get("/history?scope=all").text  # not an admin: back to Mine
        assert 'aria-current="page">Mine' in everyone and "?scope=all" not in everyone
        page = dave.get(f"/jobs/{job_id}")
        assert page.status_code == 200 and "Read-only:" in page.text and 'data-can-act="0"' in page.text
        assert dave.get(f"/api/jobs/{job_id}").json()["can_act"] is False
        assert dave.get(f"/jobs/{job_id}/report.xlsx").status_code == 200
        picked = dave.get("/api/projects/1/commits", params={"ref": "feature/login"}).json()[0]["job"]
        assert picked and picked["by"] in ("alice", "dave")

        notes = sum(len(i["notes"]) for i in fake_services.gitlab.state["issues"])
        for path in (f"/api/findings/{real['id']}/note", f"/api/findings/{real['id']}/mr",
                     f"/api/findings/{real['id']}/dismiss", f"/api/jobs/{job_id}/cancel",
                     f"/api/jobs/{job_id}/retry-failed"):
            r = dave.post(path, headers=dh)
            assert r.status_code == 403 and "Only alice can act" in r.json()["detail"], path
        assert sum(len(i["notes"]) for i in fake_services.gitlab.state["issues"]) == notes

        # Run again starts dave's own analysis, which he can act on and an admin (alice) can only view.
        own = dave.post(f"/api/jobs/{job_id}/rerun", headers=dh).json()["id"]
        assert own != job_id and dave.get(f"/api/jobs/{own}").json()["can_act"] is True
        assert client.get(f"/api/jobs/{own}").json()["can_act"] is False
        r = client.post(f"/api/jobs/{own}/cancel", headers=h)
        assert r.status_code == 403 and "Only dave can act" in r.json()["detail"]
        assert dave.post(f"/api/jobs/{own}/cancel", headers=dh).json()["status"] == "cancelled"
        assert f"/jobs/{own}\"" in client.get("/history?scope=all").text

        # Below Developer in GitLab: the analysis is hidden, as GitLab hides the project's vulnerabilities.
        fake_services.ACCESS[fake_services.BOB_TOKEN] = 20
        main._access.clear()
        try:
            assert dave.get(f"/api/jobs/{job_id}").status_code == 404
            assert dave.get(f"/jobs/{job_id}/report.xlsx").status_code == 404
            assert f"/jobs/{job_id}\"" not in dave.get("/history?scope=shared").text
            assert dave.post(f"/api/findings/{real['id']}/note", headers=dh).status_code == 404
            picked = dave.get("/api/projects/1/commits", params={"ref": "feature/login"}).json()[0]["job"]
            assert picked is None or picked["by"] == "dave"  # only his own analyses are shown
        finally:
            fake_services.ACCESS[fake_services.BOB_TOKEN] = 30
            main._access.clear()


def test_thinking_mode_per_analysis(client, fake_services, monkeypatch):
    """Thinking is chosen per analysis, sent to vLLM, and earlier results are only reused within the same mode."""
    from io import BytesIO
    from openpyxl import load_workbook
    from app.analyzer import JobRunner
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    _forget_results()
    sent = fake_services.vllm.state.thinking
    calls = _count_calls(monkeypatch)
    body = {"project_id": 1, "project_path": "grp/app", "branch": "feature/login", "sha": fake_services.SHA,
            "severities": ["high", "low"]}
    assert client.post("/api/jobs", headers=h, json={**body, "thinking": "maybe"}).status_code == 400

    sent.clear()
    off = client.post("/api/jobs", headers=h, json={**body, "thinking": "off"}).json()["id"]
    assert _claim_next() == off
    JobRunner(off).run()
    job = client.get(f"/api/jobs/{off}").json()
    assert job["thinking"] == "off" and sent and set(sent) == {False}
    for f in job["findings"]:
        assert f["status"] == "done" and f["duration_s"] > 0 and f["completion_tokens"] >= 120
    assert ">Off<" in client.get("/history").text and "Run with thinking on" in client.get(f"/jobs/{off}").text

    # The same commit with thinking on asks the model again, and the <think> block doesn't break the JSON.
    sent.clear(); calls.clear()
    on = client.post(f"/api/jobs/{off}/rerun", headers=h, json={"thinking": "on"}).json()["id"]
    assert _claim_next() == on
    JobRunner(on).run()
    job_on = client.get(f"/api/jobs/{on}").json()
    assert job_on["thinking"] == "on" and calls and set(sent) == {True}
    assert all(f["copied_from"] is None and f["verdict"] for f in job_on["findings"])
    assert all(f["completion_tokens"] > g["completion_tokens"] for f, g in zip(job_on["findings"], job["findings"]))

    # Run again keeps the mode and copies from the analysis made with it: no model calls, no model time.
    calls.clear()
    again = client.post(f"/api/jobs/{on}/rerun", headers=h).json()["id"]
    assert _claim_next() == again
    JobRunner(again).run()
    copied = client.get(f"/api/jobs/{again}").json()
    assert copied["thinking"] == "on" and calls == []
    assert all(f["copied_from"]["job"] == on and f["completion_tokens"] is None for f in copied["findings"])

    ws = load_workbook(BytesIO(client.get(f"/jobs/{off}/report.xlsx").content))
    summary = {r[1].value: r[2].value for r in ws["Summary"].iter_rows(min_row=1) if r[1].value}
    assert summary["Model thinking"] == "Off" and "per analyzed finding (2)" in summary["Analysis time"]
    head = [c.value for c in ws["Findings"][1]]
    row = [c.value for c in ws["Findings"][2]]
    assert head[3] == "Confidence" and isinstance(row[3], float)
    assert row[head.index("Analysis time (s)")] > 0 and row[head.index("Tokens generated")] >= 120
    copied_ws = load_workbook(BytesIO(client.get(f"/jobs/{again}/report.xlsx").content))["Findings"]
    assert copied_ws.cell(2, head.index("Analysis time (s)") + 1).value in (None, "")


def test_code_budget_per_analysis(client, fake_services, monkeypatch):
    """The Admin code budget limits the code per finding, is kept per analysis, and copies stay within one budget."""
    from io import BytesIO
    from openpyxl import load_workbook
    from app.analyzer import JobRunner
    from app.llm_client import llm_config
    from app.worker import _claim_next
    h = {"X-CSRF-Token": csrf(client, "/")}
    token = csrf(client, "/admin")
    cfg = llm_config()

    def set_budget(n):
        r = client.post("/admin/llm", data={"base_url": cfg["base_url"], "model": cfg["model"], "timeout_s": "120",
                                            "structured_output": "1", "context_budget": str(n), "csrf": token})
        assert "Model settings saved" in r.text

    def run(job_id):
        assert _claim_next() == job_id
        JobRunner(job_id).run()
        return client.get(f"/api/jobs/{job_id}").json()

    _forget_results()
    calls = _count_calls(monkeypatch)
    body = {"project_id": 1, "project_path": "grp/app", "branch": "feature/login", "sha": fake_services.SHA,
            "severities": ["high", "low"]}
    set_budget(4000)
    small = client.post("/api/jobs", headers=h, json=body).json()["id"]
    job = run(small)
    assert all(f["status"] == "done" and 0 < f["context_tokens"] <= 4000 for f in job["findings"])
    assert "up to 4,000 tokens" in client.get(f"/jobs/{small}").text and ">4K<" in client.get("/history").text
    ws = load_workbook(BytesIO(client.get(f"/jobs/{small}/report.xlsx").content))["Summary"]
    assert {r[1].value: r[2].value for r in ws.iter_rows() if r[1].value}["Code per finding"] == "up to 4,000 tokens"

    # Another budget asks the model again (here the code fits both budgets, so drop the exact-prompt cache too);
    # Run again with the same budget copies.
    set_budget(50000)
    from sqlalchemy import update
    from app.db import Finding, SessionLocal
    with SessionLocal() as db:
        db.execute(update(Finding).values(cache_key=""))
        db.commit()
    calls.clear()
    big = client.post(f"/api/jobs/{small}/rerun", headers=h).json()["id"]
    job = run(big)
    assert calls and all(f["copied_from"] is None for f in job["findings"])
    calls.clear()
    again = client.post(f"/api/jobs/{big}/rerun", headers=h).json()["id"]
    job = run(again)
    assert calls == [] and all(f["copied_from"]["job"] == big for f in job["findings"])
    client.post("/admin/llm/reset", data={"csrf": token})


# ---------------------------------------------------------------- Sign in with GitLab
def _gitlab_signin(fs, account: str, c=None):
    """Signs `account` in through the fake GitLab; returns (client, callback response)."""
    from urllib.parse import parse_qs, urlparse
    from app.config import settings
    from app.main import app
    c = c or TestClient(app)
    r = c.get("/login/gitlab", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith(settings.gitlab_url + "/oauth/authorize?")
    params = {k: v[0] for k, v in parse_qs(urlparse(r.headers["location"]).query).items()}
    code = fs.authorize(params, account)
    return c, c.get("/login/gitlab/callback", params={"code": code, "state": params["state"]}, follow_redirects=False)


def _user(username):
    from sqlalchemy import select
    from app.db import SessionLocal, User
    with SessionLocal() as db:
        return db.scalar(select(User).where(User.username == username))


def test_gitlab_signin_setup(client, fake_services):
    from app.main import app
    anon = TestClient(app)
    assert "Sign in with GitLab" not in anon.get("/login").text
    assert anon.get("/login/gitlab", follow_redirects=False).status_code == 404
    token = csrf(client, "/admin")
    page = client.get("/admin").text
    assert "GitLab sign-in" in page and "http://testserver/login/gitlab/callback" in page
    r = client.post("/admin/gitlab-signin", data={"enabled": "1", "app_id": "app-id-123", "csrf": token})
    assert "Enter the Application ID and secret" in r.text
    r = client.post("/admin/gitlab-signin", data={"enabled": "1", "app_id": "app-id-123", "app_url": "not a url",
                                                  "secret": "gloas-secret-456", "csrf": token})
    assert "app address must look like" in r.text
    r = client.post("/admin/gitlab-signin", data={"enabled": "1", "app_id": "app-id-123",
                                                  "secret": "gloas-secret-456", "csrf": token})
    assert "GitLab sign-in is on" in r.text and "gloas-secret-456" not in r.text
    assert "A secret is saved" in r.text
    r = client.post("/admin/gitlab-signin", data={"enabled": "1", "app_id": "app-id-123", "csrf": token})
    assert "GitLab sign-in is on" in r.text  # a blank secret keeps the saved one
    assert "Sign in with GitLab" in anon.get("/login").text


def test_gitlab_signin_flow(client, fake_services):
    from app.analyzer import JobRunner
    from app.db import SessionLocal, User
    from app.security import hash_password
    from app.worker import _claim_next
    fs = fake_services
    c, r = _gitlab_signin(fs, "carol")
    assert r.status_code == 200 and 'http-equiv="refresh" content="0;url=/"' in r.text
    assert "vt_session" in r.headers.get("set-cookie", "")
    carol = _user("carol")
    assert carol.auth_source == "gitlab" and carol.gitlab_user_id == 3 and not carol.is_admin
    assert carol.display_name == "Carol Jones" and carol.gitlab_token_type == "oauth" and carol.gitlab_refresh_enc
    page = c.get("/settings").text
    assert "through <em>Sign in with GitLab</em>" in page and "GitLab account" in page
    # No token pasted, yet analyses work with the sign-in token.
    h = {"X-CSRF-Token": csrf(c, "/")}
    assert c.get("/api/projects").status_code == 200
    _forget_results()
    job = c.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                               "sha": fs.SHA, "severities": ["high"]}).json()["id"]
    assert _claim_next() == job
    JobRunner(job).run()
    assert c.get(f"/api/jobs/{job}").json()["status"] == "done"
    # The password form and password sign-in don't apply to GitLab accounts.
    r = c.post("/settings/password", data={"current": "x", "new": "y" * 12, "csrf": h["X-CSRF-Token"]})
    assert "You sign in with GitLab" in r.text
    anon = TestClient(client.app)
    r = anon.post("/login", data={"username": "carol", "password": "anything-at-all", "csrf": csrf(anon, "/login")})
    assert "Use Sign in with GitLab" in r.text
    # Signing in again reuses the profile.
    _gitlab_signin(fs, "carol")
    assert _user("carol").id == carol.id

    # A Windows profile with the same username is reused; a local account never is.
    with SessionLocal() as db:
        db.add(User(username="erin", password_hash="windows", auth_source="ldap"))
        db.add(User(username="localdup", password_hash=hash_password("a-long-password-1")))
        db.commit()
    erin_id = _user("erin").id
    _, r = _gitlab_signin(fs, "erin")
    erin = _user("erin")
    assert r.status_code == 200 and erin.id == erin_id and erin.auth_source == "ldap" and erin.gitlab_user_id == 5
    c2, r = _gitlab_signin(fs, "localdup")
    assert r.status_code == 400 and "local account with this username" in r.text
    assert _user("localdup").gitlab_user_id is None
    _, r = _gitlab_signin(fs, "frank")
    assert r.status_code == 400 and "account isn" in r.text and _user("frank") is None


def test_gitlab_signin_rejects_bad_state(client, fake_services):
    from urllib.parse import parse_qs, urlparse
    fs = fake_services
    c = TestClient(client.app)
    r = c.get("/login/gitlab", follow_redirects=False)
    params = {k: v[0] for k, v in parse_qs(urlparse(r.headers["location"]).query).items()}
    code = fs.authorize(params, "carol")
    r = c.get("/login/gitlab/callback", params={"code": code, "state": "forged"}, follow_redirects=False)
    assert r.status_code == 400 and "start here" in r.text
    other = TestClient(client.app)  # a browser that never started the sign-in
    r = other.get("/login/gitlab/callback", params={"code": code, "state": params["state"]}, follow_redirects=False)
    assert r.status_code == 400 and "start here" in r.text
    r = c.get("/login/gitlab/callback", params={"error": "access_denied", "error_description": "The user denied"},
              follow_redirects=False)
    assert r.status_code == 400 and "cancelled" in r.text


def test_gitlab_token_renewal(client, fake_services):
    from datetime import timedelta
    from app.analyzer import JobRunner
    from app.db import SessionLocal, User, utcnow
    from app.security import decrypt_token
    from app.worker import _claim_next
    fs = fake_services
    c, _ = _gitlab_signin(fs, "carol")
    uid = _user("carol").id

    def token():
        return decrypt_token(_user("carol").gitlab_token_enc)

    # About to expire: renewed before use.
    first = token()
    with SessionLocal() as db:
        db.get(User, uid).gitlab_token_expires_at = utcnow() + timedelta(minutes=1)
        db.commit()
    fs.oauth["token_calls"].clear()
    assert c.get("/api/projects").status_code == 200
    assert fs.oauth["token_calls"] == ["refresh_token"] and token() != first and first not in fs.USERS

    # Rejected mid-way (e.g. an analysis running for hours): renewed and the request repeated.
    h = {"X-CSRF-Token": csrf(c, "/")}
    _forget_results()
    job = c.post("/api/jobs", headers=h, json={"project_id": 1, "project_path": "grp/app", "branch": "feature/login",
                                               "sha": fs.SHA, "severities": ["high"]}).json()["id"]
    assert _claim_next() == job
    runner = JobRunner(job)
    fs.revoke(token())
    fs.oauth["token_calls"].clear()
    runner.run()
    assert c.get(f"/api/jobs/{job}").json()["status"] == "done" and fs.oauth["token_calls"] == ["refresh_token"]

    # Revoked in GitLab: the token is cleared and the user is asked to sign in again.
    fs.revoke(token())
    fs.oauth["refresh"].clear()
    r = c.get("/api/projects")
    assert r.status_code in (400, 401) and "sign in with GitLab again" in r.text
    assert _user("carol").gitlab_token_enc is None
    assert "Not connected" in c.get("/settings").text


def test_fix_mr_reads_older_records():
    from types import SimpleNamespace as NS
    from app.prompts import fix_mr
    old = NS(actions=[{"type": "fix merge request !12", "url": "u", "by": "bob", "at": "2026-01-02"}],
             status="done", copied_from=None, id=1)
    assert fix_mr(old)["mr_iid"] == 12 and fix_mr(old)["by"] == "bob"
    assert fix_mr(NS(actions=[], status="done", copied_from=None, id=2)) is None
